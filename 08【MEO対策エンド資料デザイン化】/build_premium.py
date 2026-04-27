"""
MEO対策エンド資料 - Premium Redesign v3
高品質・高級感・システム感 / NG: かわいさ
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from lxml import etree

INPUT  = r'c:/Users/shira_k13727w/OneDrive/DF/アライアンス関連/資料/デザイン調整後資料/MEO/MEO対策エンド資料.pptx'
OUTPUT = r'c:/Users/shira_k13727w/OneDrive/Desktop/クロード保管用/08【MEO対策エンド資料デザイン化】/MEO対策エンド資料_premium.pptx'

# ================================================================
#  DESIGN SYSTEM — Premium Professional
# ================================================================
# Primary palette
DARK     = RGBColor(0x08, 0x14, 0x28)   # ディープネイビー（背景）
NAVY     = RGBColor(0x0D, 0x20, 0x40)   # ネイビー（ヘッダー帯）
TEAL     = RGBColor(0x00, 0xC0, 0xBE)   # クリスプティール（アクセント）
TEAL_DK  = RGBColor(0x00, 0x7A, 0x7A)   # ダークティール
TEAL_LT  = RGBColor(0x33, 0xD4, 0xD2)   # ライトティール
GOLD     = RGBColor(0xC8, 0xA2, 0x00)   # ゴールド（プレミアムアクセント）
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHT  = RGBColor(0xF3, 0xF6, 0xFA)   # コンテンツ背景
TXT_DK   = RGBColor(0x0D, 0x1E, 0x35)   # 本文ダーク
TXT_MD   = RGBColor(0x42, 0x56, 0x6E)   # 本文ミッド
RULE     = RGBColor(0xC0, 0xD4, 0xE4)   # 区切り線


def hex_of(rgb: RGBColor) -> str:
    return f'{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}'


prs = Presentation(INPUT)
W = prs.slide_width    # 12192000 EMU
H = prs.slide_height   # 6858000  EMU
slides = list(prs.slides)

# Layout constants (EMU)
HEADER_H   = int(H * 0.195)          # ヘッダー帯の高さ
GOLD_H     = int(H * 0.009)          # ゴールドラインの高さ
SIDE_W     = int(W * 0.007)          # 左アクセントバー幅
FOOT_H     = int(H * 0.010)          # フッターラインの高さ
MARGIN_L   = int(W * 0.05)           # 左マージン


# ================================================================
#  HELPER FUNCTIONS
# ================================================================

def set_bg_solid(slide, color: RGBColor):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def set_bg_gradient(slide, hex1: str, hex2: str, angle: int = 0):
    """スライド背景にグラデーション設定"""
    bg = slide.background
    fill = bg.fill
    # XML直接操作
    bgPr = bg._element.find(qn('p:bgPr'))
    if bgPr is None:
        bgPr = parse_xml('<p:bgPr xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>')
        bg._element.append(bgPr)
    # 既存fill削除
    for child in list(bgPr):
        bgPr.remove(child)
    grad_xml = (
        f'<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<a:gsLst>'
        f'<a:gs pos="0"><a:srgbClr val="{hex1}"/></a:gs>'
        f'<a:gs pos="100000"><a:srgbClr val="{hex2}"/></a:gs>'
        f'</a:gsLst>'
        f'<a:lin ang="{angle}" scaled="0"/>'
        f'</a:gradFill>'
    )
    bgPr.append(parse_xml(grad_xml))


def remove_bg_images(slide):
    """フルスライドサイズの背景画像・図形を削除"""
    sp_tree = slide.shapes._spTree
    to_del = []
    for shape in slide.shapes:
        full = (
            shape.left <= 20000 and shape.top <= 20000
            and shape.width >= W - 200000
            and shape.height >= H - 200000
        )
        if full and not hasattr(shape, 'text_frame'):
            to_del.append(shape._element)
    for el in to_del:
        try:
            sp_tree.remove(el)
        except Exception:
            pass


def _add_rect(slide, l, t, w, h, fill_rgb: RGBColor, to_back=True):
    """矩形を追加（to_back=Trueで最背面）"""
    shape = slide.shapes.add_shape(1, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    shape.line.fill.background()
    if to_back:
        el = shape._element
        sp_tree = slide.shapes._spTree
        sp_tree.remove(el)
        sp_tree.insert(2, el)
    return shape


def _add_grad_rect(slide, l, t, w, h, hex1, hex2, angle=0, to_back=True):
    """グラデーション矩形を追加"""
    shape = slide.shapes.add_shape(1, l, t, w, h)
    shape.line.fill.background()
    sp_el = shape._element
    spPr = sp_el.find(qn('p:spPr'))
    for child in list(spPr):
        if any(child.tag.endswith(x) for x in
               ['solidFill', 'gradFill', 'noFill', 'blipFill', 'pattFill']):
            spPr.remove(child)
    grad_xml = (
        f'<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<a:gsLst>'
        f'<a:gs pos="0"><a:srgbClr val="{hex1}"/></a:gs>'
        f'<a:gs pos="100000"><a:srgbClr val="{hex2}"/></a:gs>'
        f'</a:gsLst>'
        f'<a:lin ang="{angle}" scaled="0"/>'
        f'</a:gradFill>'
    )
    spPr.insert(0, parse_xml(grad_xml))
    if to_back:
        sp_tree = slide.shapes._spTree
        sp_tree.remove(sp_el)
        sp_tree.insert(2, sp_el)
    return shape


def get_max_pt(shape):
    if not hasattr(shape, 'text_frame'):
        return 0
    mx = 0
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size and r.font.size.pt > mx:
                mx = r.font.size.pt
    return mx


def style_runs(shape, color: RGBColor, bold=None, size_pt=None):
    if not hasattr(shape, 'text_frame'):
        return
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = color
            if bold is not None:
                r.font.bold = bold
            if size_pt:
                r.font.size = Pt(size_pt)


def in_header(shape) -> bool:
    """シェイプの上端がヘッダー帯内か"""
    return shape.top < HEADER_H + int(H * 0.01)


def in_footer(shape) -> bool:
    """シェイプがフッターゾーン内か"""
    return shape.top > H - FOOT_H - int(H * 0.04)


# ================================================================
#  SLIDE DESIGN LOGIC
# ================================================================

def design_cover(slide):
    """
    表紙: ディープネイビー背景 / 左ティールバー / 右パネル / ゴールドライン
    """
    remove_bg_images(slide)
    set_bg_solid(slide, DARK)

    # 右パネル（暗めのオーバーレイで奥行き感）
    _add_grad_rect(slide,
                   int(W * 0.48), 0, int(W * 0.52), H,
                   hex_of(NAVY), hex_of(DARK), angle=0)

    # 左ティールアクセントバー（太め）
    _add_rect(slide, 0, 0, int(W * 0.055), H, TEAL)

    # ゴールドラインセパレーター（縦中央やや下）
    _add_rect(slide, int(W * 0.055), int(H * 0.68),
              W - int(W * 0.055), GOLD_H, GOLD)

    # 上部ティールライン（極細）
    _add_rect(slide, 0, 0, W, int(H * 0.005), TEAL)

    # テキスト色
    for shape in slide.shapes:
        if not hasattr(shape, 'text_frame'):
            continue
        mx = get_max_pt(shape)
        if mx >= 40:
            style_runs(shape, WHITE, bold=True)
        elif mx >= 25:
            style_runs(shape, TEAL_LT, bold=True)
        else:
            style_runs(shape, WHITE)


def design_section(slide):
    """
    セクション区切り: フルダーク / ゴールド横ライン / ティールアクセント
    """
    remove_bg_images(slide)
    set_bg_solid(slide, DARK)

    # 上部ティール太ライン
    _add_rect(slide, 0, 0, W, int(H * 0.012), TEAL)

    # 下部ゴールドライン
    _add_rect(slide, 0, H - int(H * 0.012), W, int(H * 0.012), GOLD)

    # 右側ネイビーパネル（奥行き）
    _add_rect(slide, int(W * 0.62), 0, int(W * 0.38), H, NAVY)

    # 左アクセントバー（細め）
    _add_rect(slide, 0, 0, int(W * 0.008), H, TEAL)

    # ゴールド水平ルール（中央）
    _add_rect(slide, int(W * 0.06), int(H * 0.45),
              int(W * 0.55), int(H * 0.006), GOLD)

    for shape in slide.shapes:
        if not hasattr(shape, 'text_frame'):
            continue
        mx = get_max_pt(shape)
        if mx >= 60:
            style_runs(shape, GOLD, bold=True)
        elif mx >= 35:
            style_runs(shape, WHITE, bold=True)
        else:
            style_runs(shape, TEAL_LT)


def design_content(slide):
    """
    コンテンツスライド: ライト背景 / ダークヘッダー帯 / ティールアクセント
    """
    remove_bg_images(slide)
    set_bg_solid(slide, OFF_WHT)

    # ダークヘッダー帯（グラデーション）
    _add_grad_rect(slide,
                   0, 0, W, HEADER_H,
                   hex_of(DARK), hex_of(NAVY), angle=0)

    # ヘッダー直下ゴールドライン
    _add_rect(slide, 0, HEADER_H, W, GOLD_H, GOLD)

    # 左側ティールアクセントバー（ヘッダー以下のみ）
    _add_rect(slide,
              0, HEADER_H + GOLD_H,
              SIDE_W, H - HEADER_H - GOLD_H, TEAL)

    # 下部ティールフッターライン
    _add_rect(slide, 0, H - FOOT_H, W, FOOT_H, TEAL)

    # テキスト色設定
    for shape in slide.shapes:
        if not hasattr(shape, 'text_frame'):
            continue
        mx = get_max_pt(shape)
        if in_header(shape):
            # ヘッダー内 → 白太字
            style_runs(shape, WHITE, bold=True)
        else:
            # ボディ内 → サイズ別
            if mx >= 38:
                style_runs(shape, TXT_DK, bold=True)
            elif mx >= 22:
                style_runs(shape, TEAL_DK, bold=True)
            elif mx >= 15:
                style_runs(shape, TXT_DK)
            else:
                style_runs(shape, TXT_MD)


# ================================================================
#  MAIN PROCESSING
# ================================================================

COVER_IDX   = {0}
SECTION_IDX = {1}

for idx, slide in enumerate(slides):
    print(f'  Processing slide {idx + 1}/{len(slides)}...', end='')
    if idx in COVER_IDX:
        design_cover(slide)
    elif idx in SECTION_IDX:
        design_section(slide)
    else:
        design_content(slide)
    print(' done')

prs.save(OUTPUT)
print(f'\n保存完了: {OUTPUT}')
