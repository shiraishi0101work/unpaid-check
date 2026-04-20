"""
未入金チェックシステム — Webアプリ メイン

起動: python app.py
アクセス: http://localhost:5000
"""

import os
import json
import threading
from datetime import datetime, date
from functools import wraps
from flask import (
    Flask, render_template, redirect, request,
    jsonify, send_file, url_for, flash, session
)
from werkzeug.security import generate_password_hash, check_password_hash
from dotenv import load_dotenv

import auth
from mf_invoice    import get_all_billings, UNPAID_STATUSES, PAID_STATUSES
from reporter      import generate_report
from scheduler     import init_scheduler

load_dotenv()

app = Flask(__name__)
app.secret_key = os.getenv("FLASK_SECRET_KEY", "mf-unpaid-check-secret-2026")

DATA_DIR  = os.path.join(os.path.dirname(__file__), "data")
AUTH_FILE = os.path.join(DATA_DIR, "auth.json")
os.makedirs(DATA_DIR, exist_ok=True)

# ---- 会社IDごとのパスヘルパー ----
def get_company_id() -> str:
    return session.get("company_id", "c_main")

def _cdir(cid: str = None) -> str:
    """会社のデータディレクトリを返す（なければ作成）"""
    d = os.path.join(DATA_DIR, cid or get_company_id())
    os.makedirs(d, exist_ok=True)
    return d

def _hdir(cid: str = None) -> str:
    d = os.path.join(_cdir(cid), "history")
    os.makedirs(d, exist_ok=True)
    return d

# ========================================
# 認証管理
# ========================================

def load_auth() -> dict | None:
    if not os.path.exists(AUTH_FILE):
        return None
    with open(AUTH_FILE, encoding="utf-8") as f:
        return json.load(f)

def save_auth(email: str, password: str):
    data = {
        "email":         email.strip().lower(),
        "password_hash": generate_password_hash(password),
    }
    with open(AUTH_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get("logged_in"):
            return redirect(url_for("login", next=request.path))
        return f(*args, **kwargs)
    return decorated

@app.context_processor
def inject_company_name():
    s = load_settings()
    return {"company_name": s.get("company_name", "")}

@app.before_request
def check_auth():
    public = ("/login", "/setup", "/static/")
    if not any(request.path.startswith(p) for p in public):
        if not session.get("logged_in"):
            auth_data = load_auth()
            if auth_data is None:
                # 未設定 → セットアップへ
                if request.path != "/setup":
                    return redirect(url_for("setup"))
            else:
                return redirect(url_for("login", next=request.path))

@app.route("/login", methods=["GET", "POST"])
def login():
    if session.get("logged_in"):
        return redirect(url_for("dashboard"))
    auth_data = load_auth()
    if auth_data is None:
        return redirect(url_for("setup"))
    error = None
    if request.method == "POST":
        email    = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        if email == auth_data["email"] and check_password_hash(auth_data["password_hash"], password):
            session["logged_in"] = True
            session["email"]     = email
            next_url = request.form.get("next") or url_for("dashboard")
            return redirect(next_url)
        error = "メールアドレスまたはパスワードが違います"
    return render_template("login.html", error=error, next=request.args.get("next", ""))

@app.route("/logout", methods=["POST"])
def logout():
    session.clear()
    return redirect(url_for("login"))

@app.route("/setup", methods=["GET", "POST"])
def setup():
    if load_auth() is not None and not session.get("logged_in"):
        return redirect(url_for("login"))
    error = None
    if request.method == "POST":
        email    = request.form.get("email", "").strip()
        password = request.form.get("password", "")
        confirm  = request.form.get("confirm", "")
        if not email or "@" not in email:
            error = "有効なメールアドレスを入力してください"
        elif len(password) < 6:
            error = "パスワードは6文字以上にしてください"
        elif password != confirm:
            error = "パスワードが一致しません"
        else:
            save_auth(email, password)
            session["logged_in"] = True
            session["email"]     = email.lower()
            flash("アカウントを設定しました", "success")
            return redirect(url_for("dashboard"))
    return render_template("login.html", setup_mode=True, error=error, next="")

# ========================================
# 設定管理
# ========================================

DEFAULT_SETTINGS = {
    "fiscal_year_start_month": 4,   # 4月始まり
    "company_name": "",
}

def load_settings() -> dict:
    if not os.path.exists(SETTINGS_FILE):
        return dict(DEFAULT_SETTINGS)
    with open(SETTINGS_FILE, encoding="utf-8") as f:
        s = json.load(f)
    return {**DEFAULT_SETTINGS, **s}

def save_settings(settings: dict):
    with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
        json.dump(settings, f, ensure_ascii=False, indent=2)

# ========================================
# 月次履歴管理
# ========================================

def load_history_list() -> list:
    """保存済み月次履歴の一覧（新しい順）を返す"""
    files = sorted(
        [f for f in os.listdir(HISTORY_DIR) if f.endswith(".json")],
        reverse=True,
    )
    result = []
    for f in files:
        try:
            with open(os.path.join(HISTORY_DIR, f), encoding="utf-8") as fp:
                result.append(json.load(fp))
        except Exception:
            pass
    return result

def load_history(year_month: str) -> dict | None:
    path = os.path.join(HISTORY_DIR, f"{year_month}.json")
    if not os.path.exists(path):
        return None
    with open(path, encoding="utf-8") as f:
        return json.load(f)

def save_history(year_month: str, data: dict):
    path = os.path.join(HISTORY_DIR, f"{year_month}.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

def get_prev_year_month() -> str:
    """前月の YYYY-MM 文字列を返す"""
    from datetime import timedelta
    first = date.today().replace(day=1)
    return (first - timedelta(days=1)).strftime("%Y-%m")

def get_current_year_month() -> str:
    return date.today().strftime("%Y-%m")

def _build_history_snapshot(year_month: str) -> dict | None:
    """現在の last_result + meta から月次スナップショットを構築する"""
    result = _load_result()
    if not result:
        return None
    meta = load_meta()

    def apply_meta(lst):
        for b in lst:
            bid = b.get("請求書ID", "")
            m = meta.get(bid, {})
            b["担当者"]          = m.get("担当者", b.get("担当者", ""))
            b["メモ_手動"]       = m.get("メモ", b.get("メモ_手動", ""))
            b["着金日_手動"]     = m.get("着金日_手動", b.get("着金日_手動", ""))
            b["着金金額_手動"]   = m.get("着金金額_手動", b.get("着金金額_手動", 0))
            b["入金予定日_手動"] = m.get("入金予定日_手動", b.get("入金予定日_手動", ""))

    seikyu_raw   = result.get("seikyu_kakunin", [])
    seikyu_paid  = [b for b in seikyu_raw if meta.get(b.get("請求書ID",""), {}).get("着金日_手動")]
    seikyu_unpaid = [b for b in seikyu_raw if not meta.get(b.get("請求書ID",""), {}).get("着金日_手動")]

    unpaid_all = (result.get("overdue", []) + result.get("upcoming", []) +
                  result.get("unpaid", []) + seikyu_unpaid)
    paid_all   = result.get("paid", []) + seikyu_paid

    apply_meta(unpaid_all)
    apply_meta(paid_all)

    return {
        "year_month":   year_month,
        "confirmed_at": datetime.now().isoformat(),
        "summary": {
            "unpaid_count": len(unpaid_all),
            "unpaid_total": int(sum(b.get("請求金額", 0) for b in unpaid_all)),
            "paid_count":   len(paid_all),
            "paid_total":   int(sum(b.get("請求金額", 0) for b in paid_all)),
        },
        "unpaid": _clean_items(unpaid_all),
        "paid":   _clean_items(paid_all),
    }

def _generate_history_excel(history: dict) -> str:
    """月次履歴のExcelを生成してパスを返す"""
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

    def thin_border():
        s = Side(style="thin")
        return Border(left=s, right=s, top=s, bottom=s)

    def hdr(ws, r, c, val, bg, fc="FFFFFF"):
        cell = ws.cell(row=r, column=c, value=val)
        cell.fill = PatternFill("solid", fgColor=bg)
        cell.font = Font(bold=True, color=fc)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border()
        return cell

    def dat(ws, r, c, val, bg=None, fmt=None, bold=False):
        cell = ws.cell(row=r, column=c, value=val)
        if bg:
            cell.fill = PatternFill("solid", fgColor=bg)
        if fmt:
            cell.number_format = fmt
        cell.font = Font(bold=bold)
        cell.border = thin_border()
        return cell

    ym   = history["year_month"]
    summ = history["summary"]
    wb   = openpyxl.Workbook()

    # ---- サマリーシート ----
    ws = wb.active
    ws.title = "サマリー"
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 20

    hdr(ws, 1, 1, f"{ym} 月次確定レポート", "1F3864")
    ws.merge_cells("A1:B1")
    rows = [
        ("未入金件数",   summ["unpaid_count"]),
        ("未入金総額",   summ["unpaid_total"]),
        ("入金確認済み件数", summ["paid_count"]),
        ("入金確認済み総額", summ["paid_total"]),
        ("確定日時",     history.get("confirmed_at", "")[:19].replace("T", " ")),
    ]
    for i, (label, val) in enumerate(rows, start=2):
        dat(ws, i, 1, label, bold=True)
        cell = dat(ws, i, 2, val)
        if isinstance(val, int):
            cell.number_format = '#,##0'

    # ---- 未入金シート ----
    ws2 = wb.create_sheet("未入金")
    cols = ["取引先名", "請求番号", "件名", "担当者", "請求日", "支払期限", "請求金額", "メモ"]
    for c, col in enumerate(cols, 1):
        hdr(ws2, 1, c, col, "C00000")
        ws2.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 18
    ws2.column_dimensions["C"].width = 30
    ws2.column_dimensions["H"].width = 25

    for r, b in enumerate(history.get("unpaid", []), start=2):
        bg = "FCE4D6"
        dat(ws2, r, 1, b.get("取引先名", ""),  bg)
        dat(ws2, r, 2, b.get("請求番号", ""),  bg)
        dat(ws2, r, 3, b.get("件名", ""),      bg)
        dat(ws2, r, 4, b.get("担当者", ""),    bg)
        dat(ws2, r, 5, b.get("請求日", ""),    bg)
        dat(ws2, r, 6, b.get("支払期限", ""),  bg)
        dat(ws2, r, 7, b.get("請求金額", 0),   bg, '#,##0')
        dat(ws2, r, 8, b.get("メモ_手動", ""), bg)

    # ---- 入金確認済みシート ----
    ws3 = wb.create_sheet("入金確認済み")
    cols3 = ["取引先名", "請求番号", "件名", "担当者", "請求日", "支払期限", "請求金額", "着金日", "着金金額", "メモ"]
    for c, col in enumerate(cols3, 1):
        hdr(ws3, 1, c, col, "375623")
        ws3.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 16
    ws3.column_dimensions["C"].width = 30

    for r, b in enumerate(history.get("paid", []), start=2):
        bg = "E2EFDA"
        dat(ws3, r, 1, b.get("取引先名", ""),       bg)
        dat(ws3, r, 2, b.get("請求番号", ""),       bg)
        dat(ws3, r, 3, b.get("件名", ""),           bg)
        dat(ws3, r, 4, b.get("担当者", ""),         bg)
        dat(ws3, r, 5, b.get("請求日", ""),         bg)
        dat(ws3, r, 6, b.get("支払期限", ""),       bg)
        dat(ws3, r, 7, b.get("請求金額", 0),        bg, '#,##0')
        dat(ws3, r, 8, b.get("着金日_手動", ""),    bg)
        dat(ws3, r, 9, b.get("着金金額_手動", 0),   bg, '#,##0')
        dat(ws3, r,10, b.get("メモ_手動", ""),      bg)

    output_dir = os.path.join(os.path.dirname(__file__), "output")
    os.makedirs(output_dir, exist_ok=True)
    path = os.path.join(output_dir, f"月次履歴_{ym}.xlsx")
    wb.save(path)
    return path

# ========================================
# 請求書メタデータ管理（担当者・メモ・手動着金）
# ========================================

def load_meta() -> dict:
    if not os.path.exists(META_FILE):
        return {}
    with open(META_FILE, encoding="utf-8") as f:
        return json.load(f)

def save_meta(meta: dict):
    with open(META_FILE, "w", encoding="utf-8") as f:
        json.dump(meta, f, ensure_ascii=False, indent=2)


def load_company_tantou() -> dict:
    """取引先名 → 担当者 のマッピングを返す"""
    if not os.path.exists(COMPANY_TANTOU_FILE):
        return {}
    with open(COMPANY_TANTOU_FILE, encoding="utf-8") as f:
        return json.load(f)


def save_company_tantou(mapping: dict):
    with open(COMPANY_TANTOU_FILE, "w", encoding="utf-8") as f:
        json.dump(mapping, f, ensure_ascii=False, indent=2)


def apply_auto_tantou(billings: list, meta: dict) -> list:
    """
    担当者が未設定の請求書に、同じ取引先名で登録済みの担当者を自動セットする。
    meta（請求書IDベース）→ 取引先名ベースのマップも参照。
    """
    company_tantou = load_company_tantou()
    for b in billings:
        if b.get("担当者"):
            continue  # すでに設定済みはスキップ
        company = b.get("取引先名", "").strip()
        if not company:
            continue
        # 取引先名マップから自動セット
        tantou = company_tantou.get(company, "")
        if tantou:
            b["担当者"] = tantou
    return billings

# ========================================
# 年度ユーティリティ
# ========================================

def get_fiscal_year_range(year: int, start_month: int) -> tuple[str, str]:
    """年度の開始日・終了日を返す (YYYY-MM-DD)"""
    from_date = f"{year}-{start_month:02d}-01"
    end_year  = year + 1 if start_month > 1 else year
    end_month = start_month - 1 if start_month > 1 else 12
    import calendar
    last_day = calendar.monthrange(end_year, end_month)[1]
    to_date = f"{end_year}-{end_month:02d}-{last_day:02d}"
    return from_date, to_date

def get_current_fiscal_year(start_month: int) -> int:
    """今日の日付から現在の事業年度（開始年）を返す"""
    today = date.today()
    if today.month >= start_month:
        return today.year
    return today.year - 1

# ========================================
# 実行状態管理
# ========================================
_status = {
    "running":   False,
    "message":   "",
    "progress":  0,
    "error":     None,
    "last_run":  None,
    "last_file": None,
}
_status_lock = threading.Lock()

def _update_status(**kwargs):
    with _status_lock:
        _status.update(kwargs)

def load_deleted_drafts() -> set:
    """下書きとして削除済みの請求書IDセットを返す"""
    if not os.path.exists(DELETED_DRAFTS_FILE):
        return set()
    with open(DELETED_DRAFTS_FILE, encoding="utf-8") as f:
        return set(json.load(f))

def save_deleted_drafts(ids: set):
    with open(DELETED_DRAFTS_FILE, "w", encoding="utf-8") as f:
        json.dump(list(ids), f, ensure_ascii=False)

def load_deleted_seikyu() -> set:
    """未請求扱いとして除外された請求書IDセットを返す"""
    if not os.path.exists(DELETED_SEIKYU_FILE):
        return set()
    with open(DELETED_SEIKYU_FILE, encoding="utf-8") as f:
        return set(json.load(f))

def save_deleted_seikyu(ids: set):
    with open(DELETED_SEIKYU_FILE, "w", encoding="utf-8") as f:
        json.dump(list(ids), f, ensure_ascii=False)

def _is_delivered(b: dict) -> bool:
    """顧客への送付が確認されている（DL・メール・郵送のいずれか）"""
    return (
        bool(b.get("DL済み", False)) or
        b.get("メール状態") != "未送信" or
        b.get("送付状態") != "未郵送"
    )

def _is_draft_invoice(b: dict) -> bool:
    """下書き（顧客未着）: DL・メール・郵送のいずれも行われていない"""
    return not _is_delivered(b)


def _run_check(invoice_from=None, invoice_to=None, due_from=None, due_to=None):
    """MF請求書ステータス＋手動確認のみで未入金を判定する（MF会計照合なし）"""
    try:
        _update_status(running=True, error=None, progress=5, message="請求書データを取得中...")
        try:
            auth.get_access_token("invoice")
        except Exception as e:
            raise RuntimeError(f"MFクラウド請求書の認証が必要です: {e}")

        all_billings = get_all_billings()
        if not all_billings:
            raise RuntimeError("請求書データが0件でした。APIの応答を確認してください。")

        print(f"[取得] 請求書 {len(all_billings)}件")
        _update_status(progress=50, message=f"請求書 {len(all_billings)}件 取得完了。分類中...")

        today_str = datetime.now().strftime("%Y-%m-%d")
        meta = load_meta()

        # メタデータをマージ
        for b in all_billings:
            bid = b.get("請求書ID", "")
            m = meta.get(bid, {})
            b["担当者"]        = m.get("担当者", "")
            b["メモ_手動"]     = m.get("メモ", "")
            b["着金日_手動"]   = m.get("着金日_手動", "")
            b["着金金額_手動"] = m.get("着金金額_手動", 0)

        all_billings = apply_auto_tantou(all_billings, meta)

        # 日付フィルタ（支払期限ベース、未設定は請求日にフォールバック）
        def in_range(b):
            d = b.get("支払期限", "") or b.get("請求日", "")
            f = due_from or invoice_from
            t = due_to   or invoice_to
            if f and d and d < f:
                return False
            if t and d and d > t:
                return False
            return True

        filtered = [b for b in all_billings if in_range(b)]

        # 請求金額0円を除外（テスト・内部用の無効請求書）
        filtered = [b for b in filtered if (b.get("請求金額") or 0) > 0]

        # 分類
        # ① MF請求書が入金済み → 入金確認済み（MF）
        # ② MF未入金 かつ 手動着金日あり → 入金確認済み（手動）
        # ③ MF未入金 かつ 手動着金日なし → 未入金
        paid_mf     = [b for b in filtered if b["入金状態"] in PAID_STATUSES]
        paid_manual = [b for b in filtered
                       if b["入金状態"] in UNPAID_STATUSES and b.get("着金日_手動")]
        unpaid      = [b for b in filtered
                       if b["入金状態"] in UNPAID_STATUSES and not b.get("着金日_手動")]

        all_paid = paid_mf + paid_manual

        # 未入金を期限超過/期日前に分類
        all_overdue = [b for b in unpaid if b.get("支払期限") and b["支払期限"] < today_str]
        upcoming    = sorted(
            [b for b in unpaid if not b.get("支払期限") or b["支払期限"] >= today_str],
            key=lambda x: x.get("支払期限", ""))

        # 期限超過を4分類
        # ① メモ入力あり → 一時保管（毎月要確認）--- 他のタブには出さない
        # ② MFで「未入金」明示 → 未収
        # ③ MFで「未設定」かつ送付済み → 請求済みか確認
        # ④ MFで「未設定」かつ未送付 → 下書き（期限超過）
        deleted_draft_ids  = load_deleted_drafts()
        deleted_seikyu_ids = load_deleted_seikyu()

        # 一時保管：メモあり（支払済みは除く）
        ichiji_hokan = sorted(
            [b for b in all_overdue if b.get("メモ_手動")],
            key=lambda x: x.get("支払期限", ""))
        ichiji_ids = {b["請求書ID"] for b in ichiji_hokan}

        overdue = sorted(
            [b for b in all_overdue
             if b.get("入金状態") == "未入金" and b["請求書ID"] not in ichiji_ids],
            key=lambda x: x.get("支払期限", ""))
        seikyu_kakunin = sorted(
            [b for b in all_overdue
             if b.get("入金状態") == "未設定" and _is_delivered(b)
             and b["請求書ID"] not in ichiji_ids
             and b.get("請求書ID") not in deleted_seikyu_ids],
            key=lambda x: x.get("支払期限", ""))
        draft_overdue = sorted(
            [b for b in all_overdue
             if b.get("入金状態") == "未設定" and _is_draft_invoice(b)
             and b["請求書ID"] not in ichiji_ids
             and b.get("請求書ID") not in deleted_draft_ids],
            key=lambda x: x.get("支払期限", ""))

        result = {
            "paid":            all_paid,
            "unpaid":          unpaid,
            "overdue":         overdue,
            "seikyu_kakunin":  seikyu_kakunin,
            "draft_overdue":   draft_overdue,
            "ichiji_hokan":    ichiji_hokan,
            "upcoming":        upcoming,
            "total":           len(all_paid) + len(unpaid),
            "invoice_from":    invoice_from or "",
            "invoice_to":      invoice_to   or "",
        }

        print(f"[分類] 未入金={len(unpaid)}件（"
              f"未収:{len(overdue)} 請求済確認:{len(seikyu_kakunin)} 下書き:{len(draft_overdue)} "
              f"一時保管:{len(ichiji_hokan)} 期日前:{len(upcoming)}）"
              f" 入金済み={len(all_paid)}件（MF:{len(paid_mf)} 手動:{len(paid_manual)}）")

        _update_status(progress=85, message="Excelレポートを生成中...")
        output_path = generate_report(result)
        _save_result(result, output_path)

        _update_status(
            running=False, progress=100, message="完了",
            last_run=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            last_file=output_path, error=None,
        )

    except Exception as e:
        import traceback
        traceback.print_exc()
        _update_status(running=False, error=str(e), message="エラーが発生しました", progress=0)


def _clean_items(items: list) -> list:
    return [{k: v for k, v in item.items() if k != "_raw"} for item in items]


def _save_result(result: dict, output_path: str):
    data = {
        "saved_at":         datetime.now().isoformat(),
        "output_path":      output_path,
        "unpaid":           _clean_items(result["unpaid"]),
        "overdue":          _clean_items(result["overdue"]),
        "seikyu_kakunin":   _clean_items(result.get("seikyu_kakunin", [])),
        "draft_overdue":    _clean_items(result.get("draft_overdue", [])),
        "ichiji_hokan":     _clean_items(result.get("ichiji_hokan", [])),
        "upcoming":         _clean_items(result["upcoming"]),
        "paid":             _clean_items(result["paid"]),
        "total":            result["total"],
        "invoice_from":     result.get("invoice_from", ""),
        "invoice_to":       result.get("invoice_to", ""),
    }
    with open(RESULT_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def _load_result() -> dict | None:
    if not os.path.exists(RESULT_FILE):
        return None
    with open(RESULT_FILE, encoding="utf-8") as f:
        return json.load(f)

# ========================================
# ルーティング — ダッシュボード
# ========================================

@app.route("/")
def dashboard():
    settings        = load_settings()
    auth_invoice    = auth.is_authenticated("invoice")
    result          = _load_result()
    meta            = load_meta()
    today_str       = datetime.now().strftime("%Y-%m-%d")

    # 先月未確定ポップアップ判定（月初7日以内）
    prev_ym = get_prev_year_month()
    show_history_popup = (
        date.today().day <= 7 and
        load_history(prev_ym) is None and
        settings.get("history_popup_dismissed") != prev_ym
    )

    start_month = settings["fiscal_year_start_month"]
    # URL パラメータで年度切替
    fiscal_year = request.args.get("year", type=int)
    if fiscal_year is None:
        fiscal_year = get_current_fiscal_year(start_month)

    prev_year = fiscal_year - 1
    next_year = fiscal_year + 1
    fy_from, fy_to = get_fiscal_year_range(fiscal_year, start_month)

    # 年度フィルタをかけた一覧
    def in_fy(b):
        d = b.get("請求日") or b.get("支払期限") or ""
        return fy_from <= d <= fy_to

    if result:
        # メタデータを最新にマージ＋担当者自動振り分け
        company_tantou = load_company_tantou()
        for lst in [result.get("unpaid",[]), result.get("overdue",[]),
                    result.get("upcoming",[]), result.get("paid",[]),
                    result.get("seikyu_kakunin",[]), result.get("ichiji_hokan",[]),
                    result.get("draft_overdue",[])]:
            for b in lst:
                bid = b.get("請求書ID","")
                m = meta.get(bid, {})
                b["担当者"]          = m.get("担当者", b.get("担当者",""))
                b["メモ_手動"]       = m.get("メモ", b.get("メモ_手動",""))
                b["着金日_手動"]     = m.get("着金日_手動", b.get("着金日_手動",""))
                b["着金金額_手動"]   = m.get("着金金額_手動", b.get("着金金額_手動", 0))
                b["入金予定日_手動"] = m.get("入金予定日_手動", b.get("入金予定日_手動", ""))
                if not b.get("担当者"):
                    company = b.get("取引先名", "").strip()
                    if company and company in company_tantou:
                        b["担当者"] = company_tantou[company]

        # 請求済みか確認タブで着金日_手動が入力されたものは入金確認済みへ移動
        seikyu_raw     = result.get("seikyu_kakunin", [])
        seikyu_moved   = [b for b in seikyu_raw if b.get("着金日_手動")]
        seikyu_kakunin = [b for b in seikyu_raw if not b.get("着金日_手動")]

        paid_all      = result.get("paid", []) + seikyu_moved
        overdue       = result.get("overdue", [])
        draft_overdue = result.get("draft_overdue", [])
        ichiji_hokan  = result.get("ichiji_hokan", [])
        upcoming      = [b for b in result.get("upcoming", []) if in_fy(b)]
        paid_fy       = [b for b in paid_all if in_fy(b)]

    else:
        overdue = seikyu_kakunin = draft_overdue = ichiji_hokan = upcoming = paid_fy = []

    return render_template(
        "dashboard.html",
        auth_invoice=auth_invoice,
        result=result,
        overdue=overdue,
        seikyu_kakunin=seikyu_kakunin,
        draft_overdue=draft_overdue,
        ichiji_hokan=ichiji_hokan,
        upcoming=upcoming,
        paid_fy=paid_fy,
        status=dict(_status),
        today=today_str,
        settings=settings,
        fiscal_year=fiscal_year,
        prev_year=prev_year,
        next_year=next_year,
        fy_from=fy_from,
        fy_to=fy_to,
        show_history_popup=show_history_popup,
        prev_ym=prev_ym,
        current_ym=get_current_year_month(),
        history_list=load_history_list(),
    )


@app.route("/run", methods=["POST"])
def run():
    with _status_lock:
        if _status["running"]:
            return jsonify({"error": "実行中です"}), 409

    settings    = load_settings()
    start_month = settings["fiscal_year_start_month"]
    fiscal_year = request.form.get("fiscal_year", type=int)
    if fiscal_year is None:
        fiscal_year = get_current_fiscal_year(start_month)

    # 請求書発行月（YYYY-MM 形式 → YYYY-MM-01 / YYYY-MM-末日 に変換）
    invoice_from_ym = request.form.get("invoice_from_ym") or ""
    invoice_to_ym   = request.form.get("invoice_to_ym")   or ""
    def ym_to_from(ym: str) -> str | None:
        if not ym:
            return None
        return f"{ym}-01"

    def ym_to_to(ym: str) -> str | None:
        import calendar
        if not ym:
            return None
        try:
            y, m = int(ym[:4]), int(ym[5:7])
            last = calendar.monthrange(y, m)[1]
            return f"{ym}-{last:02d}"
        except Exception:
            return None

    invoice_from = ym_to_from(invoice_from_ym)
    invoice_to   = ym_to_to(invoice_to_ym)

    # 未指定なら年度から自動設定
    if not invoice_from and not invoice_to:
        invoice_from, invoice_to = get_fiscal_year_range(fiscal_year, start_month)

    thread = threading.Thread(
        target=_run_check,
        args=(invoice_from, invoice_to),
        daemon=True,
    )
    thread.start()
    return redirect(url_for("dashboard", year=fiscal_year))


@app.route("/api/status")
def api_status():
    with _status_lock:
        return jsonify(dict(_status))


@app.route("/download")
def download():
    with _status_lock:
        path = _status.get("last_file")
    if not path:
        result = _load_result()
        path = result.get("output_path") if result else None
    if not path or not os.path.exists(path):
        flash("ダウンロードできるファイルがありません。先に実行してください。")
        return redirect(url_for("dashboard"))
    return send_file(path, as_attachment=True)

# ========================================
# ルーティング — 設定
# ========================================

@app.route("/settings", methods=["GET", "POST"])
def settings_page():
    settings = load_settings()
    if request.method == "POST":
        settings["fiscal_year_start_month"] = int(request.form.get("fiscal_year_start_month", 4))
        settings["company_name"] = request.form.get("company_name", "").strip()
        save_settings(settings)
        flash("設定を保存しました", "success")
        return redirect(url_for("settings_page"))
    auth_data = load_auth() or {}
    return render_template("settings.html", settings=settings,
                           auth_invoice=auth.is_authenticated("invoice"),
                           current_email=auth_data.get("email", ""))

@app.route("/api/update-credentials", methods=["POST"])
def update_credentials():
    """メールアドレス・パスワードの変更"""
    data        = request.get_json() or {}
    new_email   = data.get("email", "").strip()
    new_password = data.get("password", "")
    confirm     = data.get("confirm", "")
    current_pw  = data.get("current_password", "")

    auth_data = load_auth()
    if not auth_data:
        return jsonify({"error": "認証情報が設定されていません"}), 400
    if not check_password_hash(auth_data["password_hash"], current_pw):
        return jsonify({"error": "現在のパスワードが違います"}), 400
    if not new_email or "@" not in new_email:
        return jsonify({"error": "有効なメールアドレスを入力してください"}), 400
    if new_password and len(new_password) < 6:
        return jsonify({"error": "パスワードは6文字以上にしてください"}), 400
    if new_password and new_password != confirm:
        return jsonify({"error": "新しいパスワードが一致しません"}), 400

    # パスワード未入力なら現在のハッシュを引き継ぐ
    if new_password:
        save_auth(new_email, new_password)
    else:
        with open(AUTH_FILE, "w", encoding="utf-8") as f:
            json.dump({"email": new_email.lower(),
                       "password_hash": auth_data["password_hash"]}, f,
                      ensure_ascii=False, indent=2)

    session["email"] = new_email.lower()
    return jsonify({"ok": True})

# ========================================
# ルーティング — 請求書メタデータ編集
# ========================================

@app.route("/api/invoice/<invoice_id>/meta", methods=["GET", "POST"])
def invoice_meta(invoice_id):
    meta = load_meta()
    if request.method == "POST":
        data = request.get_json() or {}
        meta.setdefault(invoice_id, {})
        for key in ["担当者", "メモ", "着金日_手動", "着金金額_手動", "入金予定日_手動"]:
            if key in data:
                meta[invoice_id][key] = data[key]
        save_meta(meta)

        # 担当者が設定された場合、取引先名→担当者マップも更新
        tantou = data.get("担当者", "").strip()
        if tantou:
            # 請求書IDから取引先名を逆引き
            result = _load_result()
            company_name = ""
            if result:
                all_b = (result.get("unpaid", []) + result.get("paid", [])
                         + result.get("diff", []))
                for b in all_b:
                    if b.get("請求書ID") == invoice_id:
                        company_name = b.get("取引先名", "").strip()
                        break
            if company_name:
                mapping = load_company_tantou()
                mapping[company_name] = tantou
                save_company_tantou(mapping)

        return jsonify({"ok": True})
    return jsonify(meta.get(invoice_id, {}))

# ========================================
# ルーティング — API接続テスト
# ========================================

@app.route("/api/test-connection")
def test_connection():
    results = {}

    # 請求書API
    try:
        token = auth.get_access_token("invoice")
        import requests as req
        r = req.get(
            "https://invoice.moneyforward.com/api/v3/billings",
            headers={"Authorization": f"Bearer {token}", "Accept": "application/json"},
            params={"per_page": 1, "page": 1},
            timeout=10,
        )
        if r.ok:
            total = r.json().get("pagination", {}).get("total_count", "?")
            results["invoice"] = {"ok": True, "message": f"✅ 接続OK（総件数: {total}件）"}
        else:
            results["invoice"] = {"ok": False, "message": f"❌ HTTPエラー {r.status_code}"}
    except Exception as e:
        results["invoice"] = {"ok": False, "message": f"❌ {e}"}

    return jsonify(results)



@app.route("/api/debug-check")
def debug_check():
    try:
        token = auth.get_access_token("invoice")
        billings = get_all_billings()
        from collections import Counter
        status_count = Counter(b["入金状態"] for b in billings)
        return jsonify({
            "total": len(billings),
            "token_prefix": token[:10] + "...",
            "status_counts": dict(status_count),
            "sample_due_date": billings[0].get("支払期限") if billings else None,
        })
    except Exception as e:
        import traceback
        return jsonify({"error": str(e), "trace": traceback.format_exc()}), 500

# ========================================
# OAuth認証
# ========================================

@app.route("/oauth/start/<service_name>")
def oauth_start(service_name):
    if service_name != "invoice":
        return "不正なサービス名", 400
    return redirect(auth.get_auth_url(service_name))


@app.route("/oauth/callback/<service_name>")
def oauth_callback(service_name):
    if service_name != "invoice":
        return "不正なサービス名", 400
    code  = request.args.get("code")
    state = request.args.get("state", "")
    error = request.args.get("error")
    if error:
        flash(f"認証がキャンセルされました: {error}")
        return redirect(url_for("dashboard"))
    if not code:
        flash("認可コードが取得できませんでした")
        return redirect(url_for("dashboard"))
    try:
        ok = auth.handle_callback(service_name, code, state)
        if ok:
            flash("MFクラウド請求書の認証が完了しました", "success")
        else:
            flash("認証に失敗しました（stateが一致しません）", "error")
    except Exception as e:
        flash(f"認証エラー: {e}", "error")
    return redirect(url_for("dashboard"))


@app.route("/oauth/revoke/<service_name>", methods=["POST"])
def oauth_revoke(service_name):
    if service_name != "invoice":
        return "不正なサービス名", 400
    auth.revoke_token(service_name)
    flash("MFクラウド請求書の認証を解除しました")
    return redirect(url_for("dashboard"))





@app.route("/api/invoice/<invoice_id>/delete-draft", methods=["POST"])
def delete_draft_invoice(invoice_id):
    """下書き（期限超過）タブから請求書を削除（非表示）する"""
    ids = load_deleted_drafts()
    ids.add(invoice_id)
    save_deleted_drafts(ids)

    # last_result.json からも即時除去
    result = _load_result()
    if result:
        result["draft_overdue"] = [
            b for b in result.get("draft_overdue", [])
            if b.get("請求書ID") != invoice_id
        ]
        output_path = result.get("output_path", "")
        _save_result(result, output_path)

    return jsonify({"ok": True})


@app.route("/api/invoice/mark-unmade", methods=["POST"])
def mark_unmade():
    """請求済みか確認タブの請求書を「未請求」として除外する（複数対応）"""
    data        = request.get_json() or {}
    invoice_ids = data.get("invoice_ids", [])
    if not invoice_ids:
        return jsonify({"error": "invoice_ids が必要です"}), 400

    ids = load_deleted_seikyu()
    ids.update(invoice_ids)
    save_deleted_seikyu(ids)

    # last_result.json からも即時除去
    result = _load_result()
    if result:
        id_set = set(invoice_ids)
        result["seikyu_kakunin"] = [
            b for b in result.get("seikyu_kakunin", [])
            if b.get("請求書ID") not in id_set
        ]
        _save_result(result, result.get("output_path", ""))

    return jsonify({"ok": True, "count": len(invoice_ids)})


@app.route("/api/debug-billing-raw")
def debug_billing_raw():
    """請求書APIの実際のフィールド一覧を確認するデバッグエンドポイント"""
    import requests as _req
    try:
        token = auth.get_access_token("invoice")
        r = _req.get(
            "https://invoice.moneyforward.com/api/v3/billings",
            headers={"Authorization": f"Bearer {token}", "Accept": "application/json"},
            params={"per_page": 3, "page": 1},
            timeout=30,
        )
        if r.ok:
            items = r.json().get("data", [])
            if items:
                first = items[0]
                return jsonify({
                    "all_fields": sorted(first.keys()),
                    "delivery_fields": {
                        k: first.get(k)
                        for k in ["posting_status", "posting", "email_status",
                                  "is_downloaded", "is_locked", "billing_status",
                                  "status", "draft", "payment_status"]
                        if k in first
                    },
                    "sample_3": [
                        {k: v for k, v in item.items()
                         if k not in ("billing_items", "partner", "department",
                                      "memo", "title")}
                        for item in items
                    ],
                })
        return jsonify({"error": f"HTTP {r.status_code}", "body": r.text[:500]})
    except Exception as e:
        import traceback
        return jsonify({"error": str(e), "trace": traceback.format_exc()})




# ========================================
# データクリア
# ========================================

@app.route("/api/clear-data", methods=["POST"])
def clear_data():
    if os.path.exists(RESULT_FILE):
        os.remove(RESULT_FILE)
    flash("表示データをクリアしました", "success")
    return redirect(url_for("dashboard"))

# ========================================
# 月次履歴 API・ページ
# ========================================

@app.route("/api/history/confirm", methods=["POST"])
def history_confirm():
    data = request.get_json() or {}
    year_month = data.get("year_month", "").strip()
    if not year_month:
        return jsonify({"error": "year_month が必要です"}), 400

    snapshot = _build_history_snapshot(year_month)
    if not snapshot:
        return jsonify({"error": "チェック結果がありません。先にチェックを実行してください。"}), 400

    save_history(year_month, snapshot)
    return jsonify({"ok": True, "year_month": year_month})

@app.route("/api/history/dismiss-popup", methods=["POST"])
def history_dismiss_popup():
    data = request.get_json() or {}
    year_month = data.get("year_month", "")
    s = load_settings()
    s["history_popup_dismissed"] = year_month
    save_settings(s)
    return jsonify({"ok": True})

@app.route("/history")
def history_list_page():
    histories = load_history_list()
    settings  = load_settings()
    return render_template(
        "history.html",
        histories=histories,
        detail=None,
        selected_ym=None,
        auth_invoice=auth.is_authenticated("invoice"),
        settings=settings,
    )

@app.route("/history/<year_month>")
def history_detail(year_month):
    detail = load_history(year_month)
    if not detail:
        flash(f"{year_month} の履歴が見つかりません", "error")
        return redirect(url_for("history_list_page"))
    settings = load_settings()
    return render_template(
        "history.html",
        histories=load_history_list(),
        detail=detail,
        selected_ym=year_month,
        auth_invoice=auth.is_authenticated("invoice"),
        settings=settings,
    )

@app.route("/history/<year_month>/download")
def history_download(year_month):
    detail = load_history(year_month)
    if not detail:
        flash(f"{year_month} の履歴が見つかりません", "error")
        return redirect(url_for("history_list_page"))
    path = _generate_history_excel(detail)
    return send_file(path, as_attachment=True)

# ========================================
# 入金チェックルール表
# ========================================

@app.route("/rules")
def rules_page():
    return render_template(
        "rules.html",
        auth_invoice=auth.is_authenticated("invoice"),
        settings=load_settings(),
    )

# ========================================
# 起動
# ========================================

if __name__ == "__main__":
    port  = int(os.getenv("PORT", 5000))
    debug = os.getenv("FLASK_DEBUG", "false").lower() == "true"
    init_scheduler(run_func=_run_check)
    import webbrowser, threading as _th
    _th.Timer(1.0, lambda: webbrowser.open(f"http://localhost:{port}")).start()
    print(f"\n{'='*50}")
    print(f"  未入金チェックシステム 起動中")
    print(f"  http://localhost:{port}")
    print(f"{'='*50}\n")
    app.run(host="0.0.0.0", port=port, debug=debug, use_reloader=False)
